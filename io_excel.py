"""Reads an IO ASSIGNMENT workbook and turns it into loop drawing sheets for a template set.

The workbook is read by what its columns are called, not by where they are, so the layouts of different
projects work alike (see ALIASES / GROUPS). Rows are grouped into modules: a new module starts at every yellow
highlighted row, or - when the sheet has no yellow rows - at every new MODULE NAME.

Which sheets a module gets is decided by the templates of the set that is used (see pick_templates):
every IO type / wiring / channel range the set has a template for is one sheet, e.g.
  AI  -> the 2-wire and/or 4-wire template   RTD/DI/DO -> the templates that split the channels (CH1-16 + CH17-32)
"""
import io
import re

import openpyxl

TYPES = ["AI", "AO", "RTD", "DI", "DO"]  # order of the drawing set; other types follow

# key -> header names (first match wins; "#2" = the second column with that name)
ALIASES = {
    "sr": ["SR NO"],
    "c300": ["C300 CONTROL MODULE NAME", "C300 TAG", "DCS TAG NAME"],
    "tag": ["CHANNEL NAME", "FIELD TAG"],
    "desc": ["DESCRIPTION"],
    "dcs_desc": ["DCS DESCRIPTION"],
    "equipment": ["EQUIPMENT TYPE"],
    "area": ["AREA", "SECTION"],
    "signal": ["SIGNAL", "SIGNAL POTENTIAL"],
    "wire": ["SIGNAL TYPE"],
    "controller": ["CONTROLLER NAME"],
    "link": ["LINK NO", "LINK"],
    "iom": ["IOM NO", "IOM NUM", "IOM NUMBER", "MODULE NO"],
    "channel": ["CHANNEL"],
    "module": ["MODULE NAME"],
    "iop": ["MODULE PART NO", "IOM MODEL NO", "IOP"],
    "iota": ["IOTA PART NO", "IOTA MODEL NO", "IOTA"],
    "sysgroup": ["SYSTEM TB GROUP", "IOTA TB GRP 1", "IOTA TB GROUP"],
    "sys1": ["TB1", "TB 1"],
    "sys2": ["TB2", "TB 2"],
    "sys3": ["TB3", "TB 3"],
}
# A name column followed by its terminal columns: TB NAME | TERMINAL NO | TERMINAL NO
GROUPS = {
    "tb": ["TB NAME"],
    "rtp": ["RTP NAME", "RTP NO", "RTB NAME", "RTB NO"],   # RTP / RTB = the relay terminal panel / base of a DO loop
    "jb": ["JB NAME", "JB NO"],
}
TERMINAL = re.compile(r"^(TERMINAL(NO)?\d*|TBNO\d*|RTPRTB\d*|RTB\d*|JBTERMINAL(NO)?\d*|DOTERMINAL(NO)?\d*)$")
ROW_KEYS = [*ALIASES, "iotype"]
MODULE_KEYS = ("module", "iotype", "iop", "iota", "link", "iom", "controller")
NOT_A_VALUE = {"", "NA", "N/A", "-", "--", "—"}


def clean(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return " ".join(str(v).split())


def real(v):
    return clean(v).upper() not in NOT_A_VALUE


def norm_name(v):
    """Header name without case, spaces, dots, '_' or '-': 'Module Name', 'MODULENAME' and 'MODULE_NAME' are the same."""
    return re.sub(r"[\s._\-]+", "", clean(v).upper())


def find_header(ws):
    for r, row in enumerate(ws.iter_rows(min_row=1, max_row=30, values_only=True), start=1):
        names = [norm_name(v) for v in row]
        if "MODULENAME" not in names or "CHANNEL" not in names:
            continue
        col, groups = {}, {}
        for key, options in ALIASES.items():
            for opt in options:
                name, nth = (opt.split("#") + ["1"])[:2]
                name = norm_name(name)
                hits = [i for i, n in enumerate(names) if n == name]
                if len(hits) >= int(nth):
                    col[key] = hits[int(nth) - 1]
                    break
        for key, options in GROUPS.items():
            for opt in options:
                hits = [i for i, n in enumerate(names) if n == norm_name(opt)]
                if hits:
                    i = hits[0]
                    terms, j = [], i + 1
                    while j < len(names) and TERMINAL.match(names[j]):
                        terms.append(j)
                        j += 1
                    groups[key] = (i, terms)
                    break
        # IO type: the TYPE column after LINK when there are two (old layout), else the only one
        types = [i for i, n in enumerate(names) if n in ("TYPE", "IOTYPE")]
        after = [i for i in types if i > col.get("link", -1)]
        if after or types:
            col["iotype"] = (after or types)[0]
        col["_groups"] = groups
        return r, col
    return None, None


def is_yellow_color(color):
    if color is None:
        return False
    try:
        if color.type == "rgb" and isinstance(color.rgb, str) and len(color.rgb) >= 6:
            rgb = color.rgb[-6:]
            r, g, b = int(rgb[0:2], 16), int(rgb[2:4], 16), int(rgb[4:6], 16)
            return r >= 0xC8 and g >= 0xC8 and b <= 0x99
        if color.type == "indexed":
            return color.indexed in (5, 13, 43, 51)
    except (TypeError, ValueError):
        return False
    return False


def row_is_yellow(cells, cols):
    checked = yellow = 0
    for i in cols:
        if i < len(cells):
            checked += 1
            f = cells[i].fill
            if f is not None and f.fill_type and (is_yellow_color(f.fgColor) or is_yellow_color(f.bgColor)):
                yellow += 1
    return checked > 0 and yellow >= max(2, checked // 2)


def most_common(values):
    values = [v for v in values if v]
    return max(values, key=values.count) if values else ""


def norm_type(rec):
    t = rec["iotype"].upper().replace(" ", "")
    if t and t not in ("TYPE",):
        return t
    m = rec["module"].upper()
    for name in ("RTD", "AI", "AO", "DI", "DO"):  # RTD before the 2-letter names
        if name in m:
            return name
    return t


def norm_wire(v):
    m = re.search(r"([2-4])\s*-?\s*WIRE", v.upper())
    return f"{m.group(1)} WIRE" if m else ""


def cell_value(cells, i):
    return clean(cells[i].value) if i is not None and i < len(cells) else ""


READ_ONLY_FROM = 3_000_000   # bytes of the .xlsx from which it is read in streaming mode


def read_io_excel(data: bytes, filename="", tpl=None, force_stream=False):
    """tpl: the analysed template set (tplset.load_set(...)[0]); without it the modules get no sheets."""
    try:  # a big workbook is streamed (several times faster, far less memory); cell fills are still available
        wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True, read_only=len(data) > READ_ONLY_FROM or force_stream)
    except Exception as exc:
        raise ValueError(f"not a readable .xlsx / .xlsm file ({exc})") from None
    ws = hdr = col = None
    for sheet in wb.worksheets:
        hdr, col = find_header(sheet)
        if hdr:
            ws = sheet
            break
    if not ws:
        raise ValueError("no sheet has a header row with 'MODULE NAME' and 'CHANNEL'")
    missing = [k for k in ("module", "channel", "tag", "desc") if k not in col]
    if missing:
        raise ValueError(f"columns not found in sheet '{ws.title}': {', '.join(missing)}")
    groups = col.pop("_groups")

    yellow_cols = sorted({col[k] for k in ("tag", "desc", "module", "channel") if k in col})
    rows, blank_run = [], 0
    for r, cells in enumerate(ws.iter_rows(min_row=hdr + 1), start=hdr + 1):
        rec = {k: cell_value(cells, col.get(k)) for k in ROW_KEYS}
        for g, (name_i, term_is) in groups.items():
            rec[g] = cell_value(cells, name_i)
            rec[g + "T"] = [cell_value(cells, i) for i in term_is]
        if not (rec["module"] or rec["channel"] or rec["tag"] or rec["desc"]):
            blank_run += 1
            if blank_run > 300:  # formatted but empty rows below the data
                break
            continue
        blank_run = 0
        rec["row"] = r
        rec["yellow"] = row_is_yellow(cells, yellow_cols)
        rows.append(rec)

    warnings = []
    yellow_rows = [x["row"] for x in rows if x["yellow"]]
    groups_of_rows = []
    if yellow_rows:
        mode = "yellow"
        for rec in rows:
            if rec["yellow"] or not groups_of_rows:
                groups_of_rows.append([])
            groups_of_rows[-1].append(rec)
    else:
        mode = "module"
        index = {}
        for rec in rows:
            key = rec["module"] or "(no module name)"
            if key not in index:
                index[key] = len(groups_of_rows)
                groups_of_rows.append([])
            groups_of_rows[index[key]].append(rec)

    templates = (tpl or {}).get("templates", {})
    modules = [make_module(n, g, templates) for n, g in enumerate(groups_of_rows, start=1)]
    has = {k for k in ("iop", "iota", "iotype", "link", "module", "iom", "sysgroup") if k in col} | {"chlabel"}  # 'CH5' of a one-channel sheet
    has |= {"sys"} if any(k in col for k in ("sys1", "sys2", "sys3")) else set()
    for g, (name_i, term_is) in groups.items():
        has.add(g)
    if "tb" in groups or "rtp" in groups:
        has |= {"tbname", "rtpname", "tb", "rtp"}
    if "jb" in groups:
        has |= {"jbname"}
    for m in modules:
        for sh in m["sheets"]:
            sh["texts"] = sheet_texts(m, sh, templates[sh["template"]], has)

    missing_sheets, fallback_sheets, needs = {}, {}, {}
    for m in modules:
        for label in m.pop("need"):
            n = needs.setdefault(label, {"label": label, "modules": 0, "have": label not in m["missing"]})
            n["modules"] += 1
        target = fallback_sheets if m.pop("fallback") else missing_sheets
        for w in m.pop("missing"):
            target.setdefault(w, []).append(m["module"])
    # Placeholders of the template ('LINK No: XXXX') the workbook has no column for stay as drawn
    for key, label in (("iota", "IOTA"), ("link", "LINK No"), ("iom", "IOM number"), ("module", "MODULE NAME"), ("iop", "IOP")):
        if key in has:
            continue
        left = sorted({t["header"][key]["t"] for m in modules for sh in m["sheets"] for t in [templates[sh["template"]]]
                       if key in t["header"] and re.search(r"X{2,}", t["header"][key]["t"])})
        if left:
            warnings.append(f"The workbook has no {label} column: the template text stays as drawn ({', '.join(left)}).")
    by_name = {m["module"]: m for m in modules}
    for w, mods in missing_sheets.items():
        n = sum(by_name[x]["used"] for x in mods if x in by_name)
        warnings.append(f"The template set has no {w} template - {len(mods)} module(s) ({n} used channels at most) get no sheet "
                        f"({', '.join(mods[:6])}{' …' if len(mods) > 6 else ''}). Import that template on the Templates page.")

    for w, mods in fallback_sheets.items():
        warnings.append(f"The template set has no {w} template - its channels in {len(mods)} module(s) are drawn on the template of the "
                        f"same IO type ({', '.join(mods[:6])}{' …' if len(mods) > 6 else ''}). Import a {w} template for a drawing of its own.")

    counts = {}
    for m in modules:
        c = counts.setdefault(m["type"], {"modules": 0, "sheets": 0, "used": 0, "spare": 0})
        c["modules"] += 1
        c["sheets"] += len(m["sheets"])
        c["used"] += m["used"]
        c["spare"] += m["spare"]
    return {"file": filename, "sheet": ws.title, "headerRow": hdr, "mode": mode, "dataRows": len(rows),
            "yellowRows": yellow_rows, "modules": modules, "counts": counts, "warnings": warnings, "needs": list(needs.values()),
            "set": (tpl or {}).get("id", ""), "setBuilt": (tpl or {}).get("builtAt"), "columns": sorted(k for k in col if k in ALIASES) + sorted(groups)}


def is_spare(c):
    return not c or "SPARE" in c["tag"].upper() or c["desc"].upper() in ("", "SPARE")


def sheet_texts(m, sh, t, has):
    """{text handle: full text} of one sheet: what the drawing shows for this module.

    A value the workbook has no column for stays as the template draws it; a column that is empty for this
    module / channel gives an empty text (a wrong number on a drawing is worse than none)."""
    out = {}
    values = {"iop": m["iop"], "iota": m["iota"], "iotype": m["iotype"] or m["type"], "link": m["link"], "module": m["module"],
              "iom": m["iom"], "tbname": sh["names"]["tb"], "rtpname": sh["names"]["rtp"], "sysgroup": sh["sysgroup"], "jbname": sh["jb"],
              "chlabel": str(sh["first"])}
    for key, f in t["header"].items():
        v = values.get(key, "")
        out[f["h"]] = (f["prefix"] + (v if real(v) else "")) if key in has else f["t"]
    jbf, slots = t["header"].get("jbname"), t.get("jblines", [])
    for f in slots:
        out[f["h"]] = ""
    if jbf and "jbname" in has and slots:
        lines = list_names(sh.get("jbs") or [], jbf, len(slots) + 1)
        for i, line in enumerate(lines):  # first line at the top, the last one on the line of the template
            slot = len(lines) - 1 - i
            out[(jbf if slot == 0 else slots[slot - 1])["h"]] = line
    for i, ch in enumerate(t["channels"]):
        c = sh["channels"][i]
        tag = desc = "SPARE"
        if c and c["refer"]:
            tag, desc = "-", f"REFER {c['refer']} SHEET ({c['tag']})"
        elif c:
            tag, desc = c["tag"] or "SPARE", c["desc"] or c["dcs_desc"] or "SPARE"
        out[t["tags"][i]["h"]] = t["tags"][i]["prefix"] + tag
        out[t["descs"][i]["h"]] = t["descs"][i]["prefix"] + desc
    index = {ch: i for i, ch in enumerate(t["channels"])}
    for f in t.get("jbtags", []):  # the JB of every channel box, written again at each box
        c = sh["channels"][index[f["ch"]]]
        out[f["h"]] = c["jb"] if "jbname" in has and c and real(c["jb"]) else ""
    for f in t["terms"]:
        if f["role"] not in has and not (f.get("relay") and "rtp" in has):
            out[f["h"]] = f["t"]
            continue
        c = sh["channels"][index[f["ch"]]]
        vals = c["terms"].get(f["role"], []) if c else []
        rtp = c["terms"].get("rtp", []) if c else []
        if f.get("relay") and len(rtp) >= 4 and "rtp" in has:
            # the workbook lists all four relay terminals (1+ | 1- | P1 | O1): each is drawn exactly as written
            pos = 2 + f["k"] if f["role"] == "rtp" and not f.get("derive") else f["k"]
            v = rtp[2] if f.get("derive") else rtp[pos]
            if f.get("derive"):   # the relay number 'R1' from its RTP terminal 'P1'
                n = re.search(r"\d+", v)
                v = f["derive"] + n.group() if n else ""
            out[f["h"]] = v
            continue
        v = vals[f["k"]] if f["k"] < len(vals) else ""
        if f.get("derive"):   # the relay number 'R2' from its RTP terminal 'P2'
            n = re.search(r"\d+", v)
            v = f["derive"] + n.group() if n else ""
        out[f["h"]] = v + f.get("sfx", "") if v else ""   # sfx: the polarity sign drawn behind a relay terminal ('1+')
    return out


def list_names(names, f, max_lines):
    """The JB names as a list: 'JB No. :' on top, then one JB per line. Every name is listed; only when there are more
    names than the drawing has lines, the last line takes the rest ('AJB-20 / AJB-21')."""
    lines = [f["prefix"].strip(), *names]
    if len(lines) > max_lines:
        lines = lines[:max_lines - 1] + [" / ".join(lines[max_lines - 1:])]
    return lines


def terminal_values(rec):
    """{'sys': [...], 'tb': [...], 'rtp': [...], 'jb': [...]} - '' where the workbook has no real value."""
    def vals(xs):
        out = [x if real(x) else "" for x in xs]
        while out and not out[-1]:
            out.pop()
        return out
    tb, rtp, jb = vals(rec.get("tbT", [])), vals(rec.get("rtpT", [])), vals(rec.get("jbT", []))
    sys_ = vals([rec["sys1"], rec["sys2"], rec["sys3"]])
    return {"sys": sys_, "tb": tb or rtp, "rtp": rtp or tb, "jb": jb}


def pick_templates(mtype, channels, templates):
    """(template ids, missing kinds, wiring kinds of the channels) for a module of this IO type."""
    group = sorted(((tid, t) for tid, t in templates.items() if t["type"] == mtype), key=lambda x: (x[1]["wire"], x[1]["channels"][0]))
    if not group:
        return [], [mtype or "?"], []
    if not any(t["wire"] for _, t in group):
        return [tid for tid, _ in group], [], []
    default = group[0][1]["wire"]
    kinds = sorted({norm_wire(c["wire"]) or default for c in channels.values()}) or [default]
    ids, missing = [], []
    for k in kinds:
        have = [tid for tid, t in group if t["wire"] == k]
        # a one-channel-per-sheet template is drawn for every channel, so it also serves a wiring it was not made for
        fallback = [tid for tid, t in group if t.get("perChannel")] if not have else []
        ids += [tid for tid in have + fallback if tid not in ids]
        if not have:
            missing.append(f"{mtype} {k}")
    return ids, missing, kinds


def make_module(n, group, templates):
    head = {k: most_common([rec[k] for rec in group]) for k in MODULE_KEYS}
    head["type"] = most_common([norm_type(rec) for rec in group])
    m = {"index": n, **head, "startRow": group[0]["row"], "endRow": group[-1]["row"], "rows": len(group)}
    warn = []
    names = sorted({rec["module"] for rec in group if rec["module"]})
    if len(names) > 1:
        warn.append(f"Several MODULE NAMEs in one block: {', '.join(names)}.")
    types = sorted({norm_type(rec) for rec in group if rec["iotype"]})
    if len(types) > 1:
        warn.append(f"Several IO types in one block: {', '.join(types)}.")

    # a template of one channel per sheet does not say how many channels the module has: only the ranges of the others do
    same = [t for t in templates.values() if t["type"] == m["type"] and not t.get("perChannel")]
    capacity = max((c for t in same for c in t["channels"]), default=32 if m["type"] in ("DI", "DO") else 16)
    channels = {}
    for rec in group:
        ch = int(rec["channel"]) if rec["channel"].isdigit() else None
        if ch is None or ch in channels:
            warn.append(f"Row {rec['row']}: CHANNEL '{rec['channel']}' is missing or repeated - row skipped.")
            continue
        if not 1 <= ch <= capacity:
            warn.append(f"Row {rec['row']}: CHANNEL {ch} is outside 1-{capacity} - row skipped.")
            continue
        c = {k: rec[k] for k in ("row", "tag", "c300", "desc", "dcs_desc", "wire", "sysgroup")}
        c["tb"], c["rtp"], c["jb"] = rec.get("tb", ""), rec.get("rtp", ""), rec.get("jb", "")
        c["terms"] = terminal_values(rec)
        tbT = rec.get("tbT") or rec.get("rtpT") or []
        c["tbname"] = c["tb"] if real(c["tb"]) else c["rtp"]
        c["t1"], c["t2"], c["t3"] = ((tbT + ["", "", ""])[:3])
        channels[ch] = c
    absent = [ch for ch in range(1, capacity + 1) if ch not in channels]
    if absent:
        warn.append(f"No Excel row for CH{', CH'.join(map(str, absent[:10]))}{' …' if len(absent) > 10 else ''} - shown as SPARE.")
    m["capacity"] = capacity
    m["used"] = sum(1 for c in channels.values() if not is_spare(c))
    m["spare"] = capacity - m["used"]

    ids, missing, kinds = pick_templates(m["type"], channels, templates)
    if not templates:
        warn.append("No template set is loaded.")
    elif missing and not ids:
        warn.append(f"The template set has no template for {missing[0]}.")
    elif missing and all(templates[i].get("perChannel") for i in ids):
        warn.append(f"No template for {', '.join(missing)}: those channels are drawn on the {' / '.join(templates[i]['wire'] or 'available' for i in ids)} template.")
    elif missing:
        warn.append(f"No template for {', '.join(missing)}: those channels are shown as 'REFER ... SHEET' only.")
    if len(kinds) > 1 and not all(templates[i].get("perChannel") for i in ids):
        warn.append("Module mixes " + " and ".join(kinds) + " channels: one sheet of each; the other kind is marked 'REFER ... SHEET'.")

    sheets = []
    for tid in ids:
        t = templates[tid]
        if t.get("perChannel"):
            # one sheet for every channel of the workbook, spare ones too (drawn as SPARE, so no channel number is missing
            # in the set). A channel goes to the template of its own wiring; when there is none, to this one
            # (see pick_templates)
            exact = {templates[i]["wire"] for i in ids}
            for ch in sorted(channels):
                c = channels[ch]
                k = norm_wire(c["wire"]) or kinds[0]
                if t["wire"] and k != t["wire"] and k in exact:
                    continue
                tb = c["tb"] if real(c["tb"]) else ""
                rtp = c["rtp"] if real(c["rtp"]) else ""
                jb = [c["jb"]] if real(c["jb"]) else []
                sheets.append({
                    "uid": f"{tid}@{ch}", "per": True, "template": tid, "first": ch, "last": ch,
                    "tbname": tb or rtp, "names": {"tb": tb or rtp, "rtp": rtp or tb},
                    "sysgroup": c["sysgroup"] if real(c["sysgroup"]) else "", "jb": " / ".join(jb), "jbs": jb,
                    "channels": [dict(c, ch=ch, refer="")], "used": 0 if is_spare(c) else 1,
                })
            continue
        rows_out = []
        for ch in t["channels"]:
            c = channels.get(ch)
            refer = ""
            if c and len(kinds) > 1 and not is_spare(c):  # a spare channel is spare on every sheet
                k = norm_wire(c["wire"]) or kinds[0]
                if t["wire"] and k != t["wire"]:
                    refer = k[0] + "-WIRE"
            rows_out.append(dict(c, ch=ch, refer=refer) if c else None)
        present = [c for c in rows_out if c]
        pick = lambda key: most_common([c[key] for c in present if real(c[key])])
        jb = []
        for c in present:
            if real(c["jb"]) and c["jb"] not in jb:
                jb.append(c["jb"])
        tb, rtp = pick("tb"), pick("rtp")
        sheets.append({
            "uid": tid, "per": False, "template": tid, "first": t["channels"][0], "last": t["channels"][-1],
            "tbname": tb or rtp, "names": {"tb": tb or rtp, "rtp": rtp or tb},
            "sysgroup": pick("sysgroup"), "jb": " / ".join(jb), "jbs": jb,
            "channels": rows_out,
            "used": sum(1 for c in present if not is_spare(c) and not c["refer"]),
        })
    m["sheets"] = sheets
    m["warnings"] = warn
    m["missing"] = missing
    m["fallback"] = bool(missing and ids and all(templates[i].get("perChannel") for i in ids))
    m["need"] = [f"{m['type']} {k}" for k in kinds] if kinds else [m["type"] or "?"]
    return m
